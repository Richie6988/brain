"""Espace de travail des outils portés d'iAqua : un dossier par utilisateur (fichiers, git, documents,
outils forgés), l'environnement Python dédié, le shell et les serveurs MCP.

Tout chemin est résolu dans le dossier de l'utilisateur (jamais au-dehors, liens compris). Le shell, Python
et les outils forgés exécutent du code : réservés à l'administrateur et à GUARDIAN_SHELL=1 (voir tools.py),
avec délai maximal et sorties tronquées.
"""

import json
import os
import re
import subprocess
import sys
import urllib.request
from pathlib import Path

from django.conf import settings

MAX_READ = 40_000
MAX_WRITE = 1024 * 1024
MAX_OUTPUT = 8_000
TIMEOUT = 30


class WorkspaceError(Exception):
    """Opération refusée ou impossible : le message est rendu au Gardien."""


def root(user):
    folder = Path(settings.WORKSPACE_DIR) / str(user.pk)
    folder.mkdir(parents=True, exist_ok=True)
    return folder.resolve()


def resolve(user, path):
    base = root(user)
    target = (base / (path or '.').lstrip('/')).resolve()
    if target != base and base not in target.parents:
        raise WorkspaceError(f'chemin hors de l\'espace de travail : {path!r}')
    return target


def relative(user, path):
    return str(path.relative_to(root(user)))


# --- fichiers

def read_file(user, path):
    target = resolve(user, path)
    if not target.is_file():
        raise WorkspaceError(f'fichier introuvable : {path}')
    return target.read_text('utf-8', 'replace')[:MAX_READ]


def write_file(user, path, content):
    data = (content or '').encode()
    if len(data) > MAX_WRITE:
        raise WorkspaceError('fichier trop gros (1 Mo au plus)')
    target = resolve(user, path)
    if target == root(user):
        raise WorkspaceError('chemin de fichier requis')
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return relative(user, target)


def list_files(user, path='.'):
    target = resolve(user, path)
    if not target.is_dir():
        raise WorkspaceError(f'dossier introuvable : {path}')
    return [{'name': p.name + ('/' if p.is_dir() else ''), 'size': p.stat().st_size if p.is_file() else None}
            for p in sorted(target.iterdir()) if p.name != '.git'][:200]


def edit_file(user, path, search, replace):
    text = read_file(user, path)
    count = text.count(search or '\0')
    if count != 1:
        raise WorkspaceError(f'le texte cherché apparaît {count} fois (il doit apparaître une fois)')
    return write_file(user, path, text.replace(search, replace or '', 1))


# --- exécution

def run(args, cwd, timeout=TIMEOUT, env=None, stdin=None):
    try:
        done = subprocess.run(args, cwd=cwd, capture_output=True, text=True, timeout=timeout, env=env, input=stdin)
    except subprocess.TimeoutExpired:
        raise WorkspaceError(f'délai dépassé ({timeout} s)') from None
    except OSError as e:
        raise WorkspaceError(str(e)) from None
    return {'code': done.returncode, 'stdout': done.stdout[-MAX_OUTPUT:], 'stderr': done.stderr[-MAX_OUTPUT:]}


def safe_env(user):
    return {'PATH': os.environ.get('PATH', '/usr/bin:/bin'), 'HOME': str(root(user)), 'LANG': 'C.UTF-8'}


def shell(user, command, cwd='.', timeout=TIMEOUT):
    return run(['/bin/sh', '-c', command], resolve(user, cwd), min(int(timeout or TIMEOUT), 120), safe_env(user))


def git(user, action, message='', files=None, path=''):
    base = root(user)
    identity = ['-c', 'user.name=Gardien Nodz', '-c', f'user.email=gardien@{user.pk}.nodz.local']
    if not (base / '.git').is_dir():
        run(['git', 'init', '-q'], base)
    if action == 'status':
        return run(['git', 'status', '--short', '--branch'], base)
    if action == 'diff':  # modifications non indexées puis indexées
        scope = ['--', path] if path else []
        unstaged, staged = run(['git', 'diff', *scope], base), run(['git', 'diff', '--cached', *scope], base)
        return {'code': unstaged['code'], 'stdout': (unstaged['stdout'] + staged['stdout'])[-MAX_OUTPUT:], 'stderr': unstaged['stderr']}
    if action == 'log':
        return run(['git', 'log', '--oneline', '-20'], base)
    if action == 'commit':
        if not message:
            raise WorkspaceError('message de commit requis')
        targets = [str(resolve(user, f).relative_to(base)) for f in files or []] or ['-A']
        run(['git', 'add', *targets] if targets != ['-A'] else ['git', 'add', '-A'], base)
        return run(['git', *identity, 'commit', '-q', '-m', message], base)
    raise WorkspaceError(f'action git inconnue : {action} (status, diff, log, commit)')


# --- environnement Python dédié (pyenv d'iAqua) et outils forgés

def python():
    venv = Path(settings.GUARDIAN_PYENV)
    exe = venv / 'bin' / 'python'
    if not exe.exists():
        run([sys.executable, '-m', 'venv', str(venv)], venv.parent if venv.parent.exists() else Path('/'), 120)
    return str(exe)


def pyenv(action, packages=()):
    packages = [p for p in packages or [] if re.fullmatch(r'[A-Za-z0-9_.\-\[\]=<>~!]+', p)]
    if action == 'list':
        return run([python(), '-m', 'pip', 'list', '--format=freeze'], '/', 60)
    if not packages:
        raise WorkspaceError('paquets requis')
    if action == 'install':
        return run([python(), '-m', 'pip', 'install', *packages], '/', 300)
    if action == 'remove':
        return run([python(), '-m', 'pip', 'uninstall', '-y', *packages], '/', 120)
    raise WorkspaceError(f'action inconnue : {action} (install, list, remove)')


def tool_path(user, name):
    if not re.fullmatch(r'[a-z][a-z0-9_]{2,30}', name or ''):
        raise WorkspaceError('nom d\'outil en snake_case, 3 à 31 caractères')
    return resolve(user, f'tools/{name}.py')


def run_tool(user, name, payload):
    """Un outil forgé lit son entrée JSON sur stdin et écrit son résultat sur stdout."""
    path = tool_path(user, name)
    if not path.is_file():
        raise WorkspaceError(f'outil inconnu : {name}')
    return run([python(), str(path)], root(user), TIMEOUT, safe_env(user), json.dumps(payload or {}))


# --- documents (generate_pptx, generate_docx)

def export_path(user, filename, suffix):
    name = re.sub(r'[^\w.\- ]', '', filename or 'document').strip() or 'document'
    return resolve(user, f"exports/{name if name.endswith(suffix) else name + suffix}")


# --- modèles de documents de l'utilisateur (bibliothèque du Rédacteur) : ses fichiers pptx, docx, xlsx, pdf, rangés
# dans « modeles » de son espace ; un document généré « dans son style » part de l'un d'eux.

TEMPLATE_DIR = 'modeles'
TEMPLATE_KINDS = ('.pptx', '.docx', '.xlsx', '.pdf')
TEMPLATE_MAX = 20 * 1024 * 1024


def templates(user):
    folder = resolve(user, TEMPLATE_DIR)
    if not folder.is_dir():
        return []
    return [{'name': p.stem, 'kind': p.suffix[1:], 'size': p.stat().st_size}
            for p in sorted(folder.iterdir()) if p.is_file() and p.suffix.lower() in TEMPLATE_KINDS]


def save_template(user, filename, data):
    stem, dot, ext = str(filename or '').rpartition('.')
    ext = f'.{ext.lower()}'
    name = re.sub(r'[^\w\- ]+', '', stem if dot else '').strip()[:60]
    if ext not in TEMPLATE_KINDS or not name:
        raise WorkspaceError('modèle : un fichier .pptx, .docx, .xlsx ou .pdf, avec un nom')
    if len(data) > TEMPLATE_MAX:
        raise WorkspaceError('modèle : 20 Mo au plus')
    folder = resolve(user, TEMPLATE_DIR)
    folder.mkdir(parents=True, exist_ok=True)
    for old in folder.glob(f'{name}.*'):  # un nom, un modèle : le nouveau remplace l'ancien
        old.unlink()
    (folder / f'{name}{ext}').write_bytes(data)
    return name


def template_path(user, name, kind):
    """Le fichier du modèle `name` de ce type, ou None (pas de modèle : document par défaut)."""
    if not name:
        return None
    path = resolve(user, f'{TEMPLATE_DIR}/{name}.{kind}')
    if not path.is_file():
        raise WorkspaceError(f'modèle de document introuvable : {name} ({kind})')
    return path


def delete_template(user, name):
    for path in resolve(user, TEMPLATE_DIR).glob(f'{re.sub(r"[^\w\- ]+", "", str(name))}.*'):
        path.unlink()


def pptx(user, filename, title, slides, template=None):
    from pptx import Presentation
    from pptx.util import Pt

    source = template_path(user, template, 'pptx')
    deck = Presentation(str(source)) if source else Presentation()
    if source:  # le modèle donne thème, polices et mises en page ; ses diapositives d'exemple partent
        listing = deck.slides._sldIdLst
        for slide in list(listing):
            deck.part.drop_rel(slide.rId)
            listing.remove(slide)
    layout = lambda i: deck.slide_layouts[min(i, len(deck.slide_layouts) - 1)]
    if title:
        cover = deck.slides.add_slide(layout(0))
        if cover.shapes.title is not None:
            cover.shapes.title.text = title
    for slide in slides or []:
        page = deck.slides.add_slide(layout(1))
        if page.shapes.title is not None:
            page.shapes.title.text = str(slide.get('title', ''))
        holders = [p for p in page.placeholders if p.placeholder_format.idx != 0]
        if not holders:
            continue
        body = holders[0].text_frame
        lines = slide.get('bullets') or [line for line in str(slide.get('body', '')).split('\n') if line.strip()]
        for i, line in enumerate(lines):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = str(line)
            if not source:
                paragraph.font.size = Pt(20)  # avec un modèle, ses tailles priment
    target = export_path(user, filename, '.pptx')
    target.parent.mkdir(parents=True, exist_ok=True)
    deck.save(target)
    return relative(user, target)


def docx(user, filename, title, markdown, template=None):
    from docx import Document

    source = template_path(user, template, 'docx')
    document = Document(str(source)) if source else Document()
    if source:  # styles, en-têtes et pieds de page du modèle ; son contenu d'exemple part
        body = document.element.body
        for child in list(body):
            if not child.tag.endswith('sectPr'):
                body.remove(child)
    if title:
        document.add_heading(title, 0)
    for line in (markdown or '').split('\n'):
        heading = re.match(r'(#{1,4})\s+(.*)', line)
        if heading:
            document.add_heading(heading.group(2), len(heading.group(1)))
        elif re.match(r'\s*[-*]\s+', line):
            styles = {s.name for s in document.styles}
            document.add_paragraph(re.sub(r'^\s*[-*]\s+', '', line), style='List Bullet' if 'List Bullet' in styles else None)
        elif line.strip():
            document.add_paragraph(line)
    target = export_path(user, filename, '.docx')
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return relative(user, target)


def xlsx(user, filename, title, rows, template=None):
    """Un classeur : `rows` (listes de cellules) écrites à la suite du modèle (ses en-têtes et styles gardés)."""
    from openpyxl import Workbook, load_workbook

    source = template_path(user, template, 'xlsx')
    book = load_workbook(str(source)) if source else Workbook()
    sheet = book.active
    if title and not source:
        sheet.title = str(title)[:31]
    filled = source and any(cell.value is not None for row in sheet.iter_rows() for cell in row)
    start = sheet.max_row + 1 if filled else 1  # sous les en-têtes du modèle
    for r, row in enumerate(rows or [], start):
        for c, value in enumerate(row if isinstance(row, list) else [row], 1):
            sheet.cell(row=r, column=c, value=value if isinstance(value, (int, float)) else str(value))
    target = export_path(user, filename, '.xlsx')
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return relative(user, target)


def pdf(user, filename, title, markdown, template=None):
    """Un PDF : le texte (titres #, listes -) mis en page ; avec un modèle, chaque page est posée sur sa première page
    (papier à en-tête, fond, logo)."""
    import io

    from PyPDF2 import PdfReader, PdfWriter
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()
    flow = [Paragraph(escape(title), styles['Title'])] if title else []
    for line in (markdown or '').split('\n'):
        heading = re.match(r'(#{1,3})\s+(.*)', line)
        if heading:
            flow.append(Paragraph(escape(heading.group(2)), styles[f'Heading{len(heading.group(1))}']))
        elif re.match(r'\s*[-*]\s+', line):
            flow.append(ListFlowable([Paragraph(escape(re.sub(r'^\s*[-*]\s+', '', line)), styles['BodyText'])], bulletType='bullet'))
        elif line.strip():
            flow.append(Paragraph(escape(line), styles['BodyText']))
        else:
            flow.append(Spacer(1, 6))
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, topMargin=90, bottomMargin=70).build(flow or [Spacer(1, 1)])
    content = PdfReader(io.BytesIO(buffer.getvalue()))
    writer = PdfWriter()
    source = template_path(user, template, 'pdf')
    for page in content.pages:
        if source:
            background = PdfReader(str(source)).pages[0]
            background.merge_page(page)
            page = background
        writer.add_page(page)
    target = export_path(user, filename, '.pdf')
    target.parent.mkdir(parents=True, exist_ok=True)
    with open(target, 'wb') as out:
        writer.write(out)
    return relative(user, target)


# --- serveurs MCP (configurés par l'administrateur dans MCP_SERVERS)

def mcp_servers():
    try:
        return json.loads(settings.MCP_SERVERS or '{}')
    except json.JSONDecodeError:
        raise WorkspaceError('MCP_SERVERS n\'est pas un JSON valide') from None


def _rpc(url, headers, method, params, session=None, ident=1):
    body = json.dumps({'jsonrpc': '2.0', 'id': ident, 'method': method, 'params': params or {}}).encode()
    request = urllib.request.Request(url, body, {'Content-Type': 'application/json', 'Accept': 'application/json, text/event-stream',
                                                 **headers, **({'Mcp-Session-Id': session} if session else {})})
    try:
        with urllib.request.urlopen(request, timeout=TIMEOUT) as response:
            text = response.read(2 * 1024 * 1024).decode('utf-8', 'replace')
            session = response.headers.get('Mcp-Session-Id') or session
    except OSError as e:
        raise WorkspaceError(f'serveur MCP injoignable : {e}') from None
    events = [line[5:].strip() for line in text.splitlines() if line.startswith('data:')] or [text]
    reply = json.loads(events[-1] or '{}')
    if 'error' in reply:
        raise WorkspaceError(f"MCP : {reply['error'].get('message', reply['error'])}")
    return reply.get('result', {}), session


def mcp_call(server, name, arguments=None):
    config = mcp_servers().get(server)
    if not config or not config.get('url'):
        raise WorkspaceError(f'serveur MCP inconnu : {server}')
    headers = config.get('headers', {})
    _, session = _rpc(config['url'], headers, 'initialize', {
        'protocolVersion': '2025-03-26', 'capabilities': {}, 'clientInfo': {'name': 'nodz-gardien', 'version': '1.0'}})
    if name == 'tools/list':
        result, _ = _rpc(config['url'], headers, 'tools/list', {}, session, 2)
        return [{'name': t.get('name'), 'description': t.get('description', '')[:200]} for t in result.get('tools', [])]
    result, _ = _rpc(config['url'], headers, 'tools/call', {'name': name, 'arguments': arguments or {}}, session, 2)
    texts = [c.get('text', '') for c in result.get('content', []) if c.get('type') == 'text']
    return '\n'.join(texts)[:MAX_OUTPUT] or json.dumps(result)[:MAX_OUTPUT]
