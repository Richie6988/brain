"""Texte d'un fichier envoyé (recherche, contexte de l'IA), comme les extracteurs de v1."""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)
MAX_CHARS = 200_000


def _pdf(path):
    from PyPDF2 import PdfReader

    return '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)


def _docx(path):
    from docx import Document

    return '\n'.join(p.text for p in Document(path).paragraphs)


def _xlsx(path):
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    return '\n'.join(' '.join(str(c) for c in row if c is not None) for sheet in book for row in sheet.iter_rows(values_only=True))


def _pptx(path):
    from pptx import Presentation

    return '\n'.join(shape.text for slide in Presentation(path).slides for shape in slide.shapes if hasattr(shape, 'text'))


def _plain(path):
    return Path(path).read_text(encoding='utf-8', errors='replace')


EXTRACTORS = {'.pdf': _pdf, '.docx': _docx, '.xlsx': _xlsx, '.pptx': _pptx,
              '.txt': _plain, '.md': _plain, '.csv': _plain, '.json': _plain}


def extract_text(path):
    extractor = EXTRACTORS.get(Path(path).suffix.lower())
    if extractor is None:
        return ''
    try:
        return extractor(path)[:MAX_CHARS]
    except Exception:  # fichier corrompu ou protégé : le fichier reste utilisable sans son texte
        logger.warning('extraction impossible : %s', path, exc_info=True)
        return ''
